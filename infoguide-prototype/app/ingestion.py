import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
import pdfplumber
from pdf2image import convert_from_path
import pytesseract
from pypdf import PdfReader

try:
    import fitz
except ImportError:
    fitz = None

RAW_DATA_DIR = Path("data/raw")
IMAGES_DIR = Path("data/images")

MIN_MEANINGFUL_ALNUM_CHARS = 20
PDF_TEXT_X_TOLERANCE = 2
PDF_TEXT_Y_TOLERANCE = 2

RELS_IMAGE_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

_BLIP_TAG = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_IMAGEDATA_TAG = "{urn:schemas-microsoft-com:vml}imagedata"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def save_uploaded_file(uploaded_file):
    """
    Uploaded file'ı data/raw klasörüne kaydeder.
    """
    RAW_DATA_DIR.mkdir(parents=True, exist_ok=True)

    file_path = RAW_DATA_DIR / uploaded_file.name

    with open(file_path, "wb") as f:
        f.write(uploaded_file.getbuffer())

    return file_path


def sanitize_document_stem(stem):
    sanitized = re.sub(r"[^\w\-.]+", "_", stem).strip("._")
    return sanitized or "document"


def get_document_image_dir(file_path):
    file_path = Path(file_path)
    document_stem = sanitize_document_stem(file_path.stem)
    image_dir = IMAGES_DIR / document_stem
    image_dir.mkdir(parents=True, exist_ok=True)
    return image_dir


def to_project_relative_path(path):
    path = Path(path)
    if not path.is_absolute():
        return path.as_posix()

    try:
        return path.resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return path.as_posix()


def save_image_bytes(image_bytes, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "wb") as f:
        f.write(image_bytes)
    return output_path


def format_image_marker(image_path, page_number=None, image_number=None, source=None):
    relative_path = to_project_relative_path(image_path)

    if source == "docx":
        image_index = image_number if image_number is not None else 1
        return f"[Image extracted: {relative_path} | source=docx | image={image_index}]"

    page_index = page_number if page_number is not None else 0
    image_index = image_number if image_number is not None else 1
    return (
        f"[Image extracted: {relative_path} | page={page_index} | image={image_index}]"
    )


def normalize_image_extension(ext):
    if not ext:
        return "png"

    ext = ext.lower().lstrip(".")
    if ext == "jpeg":
        return "jpg"
    if ext in {"png", "jpg", "gif", "bmp", "tif", "tiff", "webp"}:
        return ext
    return "png"


def extract_images_from_pdf(file_path):
    """
    Extract embedded PDF images per page and return page-level markers.
    """
    file_path = Path(file_path)
    image_dir = get_document_image_dir(file_path)
    page_markers = {}
    warnings = []

    if fitz is None:
        warnings.append(
            "[Warning: PyMuPDF is not installed. PDF image extraction skipped. "
            "Install PyMuPDF to enable embedded image extraction.]"
        )
        return page_markers, warnings

    doc = None
    try:
        doc = fitz.open(file_path)

        for page_index in range(len(doc)):
            page_number = page_index + 1
            page = doc[page_index]
            seen_xrefs = set()
            image_number = 0

            try:
                image_list = page.get_images(full=True)
            except Exception as e:
                warnings.append(
                    f"[Warning: Image extraction failed on page {page_number}: {e}]"
                )
                continue

            for image_info in image_list:
                try:
                    xref = image_info[0]
                    if xref in seen_xrefs:
                        continue
                    seen_xrefs.add(xref)

                    base_image = doc.extract_image(xref)
                    image_bytes = base_image.get("image")
                    if not image_bytes:
                        continue

                    ext = normalize_image_extension(base_image.get("ext"))
                    image_number += 1
                    filename = f"page_{page_number:03d}_image_{image_number:03d}.{ext}"
                    output_path = image_dir / filename
                    save_image_bytes(image_bytes, output_path)

                    marker = format_image_marker(
                        output_path,
                        page_number=page_number,
                        image_number=image_number,
                    )
                    page_markers.setdefault(page_number, []).append(marker)
                except Exception as e:
                    warnings.append(
                        f"[Warning: Image extraction failed on page {page_number}: {e}]"
                    )
    except Exception as e:
        warnings.append(f"[Warning: PDF image extraction failed: {e}]")
    finally:
        if doc is not None:
            doc.close()

    return page_markers, warnings


def _extract_docx_images_with_rid_map(file_path, image_dir):
    """
    Extract all DOCX images from the ZIP package and build per-part rId → path maps.
    Returns (rid_maps: dict[str, dict[str, Path]], warnings: list[str]).
    rid_maps keys are part paths like 'word/document.xml'.
    """
    file_path = Path(file_path)
    image_counter = 0
    media_to_path = {}
    rid_maps = {}
    warnings = []

    try:
        with zipfile.ZipFile(file_path, "r") as z:
            for name in sorted(z.namelist()):
                if name.startswith("word/media/") and not name.endswith("/"):
                    try:
                        image_bytes = z.read(name)
                        if not image_bytes:
                            continue
                        ext = normalize_image_extension(Path(name).suffix.lstrip("."))
                        image_counter += 1
                        filename = f"docx_image_{image_counter:03d}.{ext}"
                        output_path = image_dir / filename
                        save_image_bytes(image_bytes, output_path)
                        media_to_path[name[len("word/"):]] = output_path
                    except Exception as e:
                        warnings.append(f"[Warning: DOCX image save failed for {name}: {e}]")

            for name in z.namelist():
                if not (name.startswith("word/_rels/") and name.endswith(".rels")):
                    continue
                part_name = name.replace("word/_rels/", "word/")
                if part_name.endswith(".rels"):
                    part_name = part_name[:-5]
                try:
                    root = ET.fromstring(z.read(name))
                    rid_map = {}
                    for rel in root:
                        if RELS_IMAGE_TYPE not in rel.get("Type", ""):
                            continue
                        rid = rel.get("Id")
                        target = rel.get("Target", "").lstrip("../")
                        if rid and target in media_to_path:
                            rid_map[rid] = media_to_path[target]
                    if rid_map:
                        rid_maps[part_name] = rid_map
                except Exception as e:
                    warnings.append(f"[Warning: DOCX rels parse failed for {name}: {e}]")
    except Exception as e:
        warnings.append(f"[Warning: DOCX image extraction failed: {e}]")

    return rid_maps, warnings


def _get_paragraph_image_rids(paragraph):
    rids = []
    for blip in paragraph._p.iter(_BLIP_TAG):
        rid = blip.get(_R_EMBED)
        if rid:
            rids.append(rid)
    for imgdata in paragraph._p.iter(_IMAGEDATA_TAG):
        rid = imgdata.get(_R_ID)
        if rid:
            rids.append(rid)
    return rids


def clean_text(text):
    if text is None:
        return ""

    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def clean_cell(cell):
    if cell is None:
        return ""

    value = str(cell).replace("\r\n", "\n").replace("\r", "\n")
    value = value.replace("\t", " ").replace("\n", " ")
    value = re.sub(r"\s+", " ", value)
    value = value.strip()
    return value.replace("|", "\\|")


def is_meaningful_text(text):
    cleaned = clean_text(text)
    if not cleaned:
        return False

    alnum_count = sum(char.isalnum() for char in cleaned)
    return alnum_count >= MIN_MEANINGFUL_ALNUM_CHARS


def normalize_table(table):
    if not table:
        return []

    normalized_rows = []

    for row in table:
        if row is None:
            continue

        cleaned_row = [clean_cell(cell) for cell in row]
        if any(cell for cell in cleaned_row):
            normalized_rows.append(cleaned_row)

    if not normalized_rows:
        return []

    column_count = max(len(row) for row in normalized_rows)
    return [row + [""] * (column_count - len(row)) for row in normalized_rows]


def looks_like_header(row):
    non_empty = [cell for cell in row if cell]
    if len(non_empty) < 2:
        return False

    unique_cells = len({cell.lower() for cell in non_empty})
    if unique_cells < 2:
        return False

    numeric_like = sum(
        1 for cell in non_empty
        if re.fullmatch(r"[\d.,\-%+\s]+", cell)
    )
    if numeric_like >= len(non_empty):
        return False

    return True


def dedupe_consecutive_rows(rows):
    if not rows:
        return []

    deduped = [rows[0]]
    for row in rows[1:]:
        if row != deduped[-1]:
            deduped.append(row)
    return deduped


def format_markdown_table(table):
    normalized_rows = normalize_table(table)
    if not normalized_rows:
        return ""

    normalized_rows = dedupe_consecutive_rows(normalized_rows)
    column_count = len(normalized_rows[0])

    if looks_like_header(normalized_rows[0]):
        header = normalized_rows[0]
        data_rows = normalized_rows[1:]
    else:
        header = [f"Column {index}" for index in range(1, column_count + 1)]
        data_rows = normalized_rows

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * column_count) + " |",
    ]

    for row in data_rows:
        if any(cell for cell in row):
            lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def extract_pdf_page_text(pdfplumber_page, pypdf_page, page_number, table_bboxes=None):
    page_text = ""

    if pdfplumber_page is not None:
        try:
            if table_bboxes:
                def _outside_tables(obj):
                    if obj.get("object_type") != "char":
                        return True
                    cx0, cy0, cx1, cy1 = obj["x0"], obj["top"], obj["x1"], obj["bottom"]
                    for bx0, by0, bx1, by1 in table_bboxes:
                        if cx0 >= bx0 - 1 and cy0 >= by0 - 1 and cx1 <= bx1 + 1 and cy1 <= by1 + 1:
                            return False
                    return True
                filtered = pdfplumber_page.filter(_outside_tables)
                page_text = filtered.extract_text(
                    layout=True,
                    x_tolerance=PDF_TEXT_X_TOLERANCE,
                    y_tolerance=PDF_TEXT_Y_TOLERANCE,
                ) or ""
            else:
                page_text = pdfplumber_page.extract_text(
                    layout=True,
                    x_tolerance=PDF_TEXT_X_TOLERANCE,
                    y_tolerance=PDF_TEXT_Y_TOLERANCE,
                ) or ""
        except Exception:
            page_text = ""

        if is_meaningful_text(page_text):
            return clean_text(page_text), "pdfplumber"

    if pypdf_page is not None:
        try:
            fallback_text = pypdf_page.extract_text() or ""
        except Exception:
            fallback_text = ""

        if is_meaningful_text(fallback_text):
            return clean_text(fallback_text), "pypdf"

    return "", "none"


def _find_pdf_page_tables(pdfplumber_page):
    try:
        return pdfplumber_page.find_tables() or []
    except Exception:
        return []


def _get_table_bboxes(tables):
    bboxes = []
    for table in tables:
        try:
            bboxes.append(table.bbox)
        except Exception:
            pass
    return bboxes


def extract_pdf_page_tables_from_objects(tables, page_number):
    rendered_tables = []
    for table_index, table in enumerate(tables, start=1):
        try:
            extracted = table.extract()
            if not extracted:
                continue
            markdown_table = format_markdown_table(extracted)
            if not markdown_table.strip():
                continue
            rendered_tables.append(
                f"--- Page {page_number} Table {table_index} ---\n"
                f"[Table extracted from page {page_number}, table {table_index}]\n"
                f"{markdown_table}"
            )
        except Exception as e:
            rendered_tables.append(
                f"[Warning: Table extraction failed on page {page_number}: {e}]"
            )
    return rendered_tables


def run_ocr_for_page(file_path, page_number):
    images = convert_from_path(
        file_path,
        first_page=page_number,
        last_page=page_number,
    )

    ocr_text = pytesseract.image_to_string(images[0], lang="eng+tur")
    return clean_text(ocr_text)


def extract_text_from_pdf(file_path):
    """
    PDF dosyasından metin çıkarır.
    Sayfa bazlı selectable text + tablo extraction + OCR fallback + image markers.
    """
    file_path = Path(file_path)
    output_parts = []

    page_image_markers, image_warnings = extract_images_from_pdf(file_path)
    output_parts.extend(image_warnings)

    reader = PdfReader(file_path)

    with pdfplumber.open(file_path) as pdf:
        total_pages = max(len(reader.pages), len(pdf.pages))

        for page_number in range(1, total_pages + 1):
            output_parts.append(f"--- Page {page_number} ---")

            try:
                pypdf_page = (
                    reader.pages[page_number - 1]
                    if page_number <= len(reader.pages)
                    else None
                )
                pdfplumber_page = (
                    pdf.pages[page_number - 1]
                    if page_number <= len(pdf.pages)
                    else None
                )

                if pdfplumber_page is None and pypdf_page is None:
                    output_parts.append(
                        f"[Warning: Text extraction failed on page {page_number}: page not found]"
                    )
                    continue

                found_tables = []
                table_bboxes = []
                if pdfplumber_page is not None:
                    found_tables = _find_pdf_page_tables(pdfplumber_page)
                    table_bboxes = _get_table_bboxes(found_tables)

                page_text, _ = extract_pdf_page_text(
                    pdfplumber_page,
                    pypdf_page,
                    page_number,
                    table_bboxes=table_bboxes,
                )

                if page_text:
                    output_parts.append(page_text)
                else:
                    try:
                        ocr_text = run_ocr_for_page(file_path, page_number)
                        if ocr_text:
                            output_parts.append(ocr_text)
                        else:
                            output_parts.append(
                                f"[Warning: OCR failed on page {page_number}: empty OCR output]"
                            )
                    except Exception as e:
                        output_parts.append(
                            f"[Warning: OCR failed on page {page_number}: {e}]"
                        )

                if found_tables:
                    output_parts.extend(
                        extract_pdf_page_tables_from_objects(found_tables, page_number)
                    )

                output_parts.extend(page_image_markers.get(page_number, []))
            except Exception as e:
                output_parts.append(
                    f"[Warning: Text extraction failed on page {page_number}: {e}]"
                )

    return "\n\n".join(part for part in output_parts if part and part.strip())


def extract_text_from_txt(file_path):
    """
    TXT dosyasından metin çıkarır.
    """
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()

    return content.replace("\r\n", "\n").replace("\r", "\n")


def docx_table_to_rows(table):
    rows = []
    for row in table.rows:
        seen_tcs = set()
        cells = []
        for cell in row.cells:
            tc_id = id(cell._tc)
            if tc_id not in seen_tcs:
                seen_tcs.add(tc_id)
                cells.append(cell.text)
        rows.append(cells)
    return rows


def append_docx_table(text_parts, marker, table):
    text_parts.append(marker)
    try:
        markdown_table = format_markdown_table(docx_table_to_rows(table))
        if markdown_table:
            text_parts.append(markdown_table)
        else:
            text_parts.append("[Warning: Empty or invalid table]")
    except Exception as e:
        text_parts.append(f"[Warning: Table extraction failed: {e}]")


def extract_text_from_docx(file_path):
    """
    DOCX dosyasından paragraph, table, header ve footer metinlerini çıkarır.
    Görüntüler belgede göründükleri konumda satır içine yerleştirilir.
    """
    file_path = Path(file_path)
    image_dir = get_document_image_dir(file_path)
    rid_maps, image_warnings = _extract_docx_images_with_rid_map(file_path, image_dir)

    main_rid_map = rid_maps.get("word/document.xml", {})
    header_rid_map = {}
    footer_rid_map = {}
    for part_name, rid_map in rid_maps.items():
        if "header" in part_name.lower():
            header_rid_map.update(rid_map)
        elif "footer" in part_name.lower():
            footer_rid_map.update(rid_map)

    doc = Document(file_path)
    text_parts = []
    text_parts.extend(image_warnings)

    placed_rids = set()
    inline_counter = [0]

    def _place_images(paragraph, rid_map):
        for rid in _get_paragraph_image_rids(paragraph):
            if rid in rid_map and rid not in placed_rids:
                inline_counter[0] += 1
                text_parts.append(
                    format_image_marker(
                        rid_map[rid], image_number=inline_counter[0], source="docx"
                    )
                )
                placed_rids.add(rid)

    text_parts.append("--- Document Paragraphs ---")
    try:
        for paragraph in doc.paragraphs:
            paragraph_text = clean_text(paragraph.text)
            if paragraph_text:
                text_parts.append(paragraph_text)
            _place_images(paragraph, main_rid_map)
    except Exception as e:
        text_parts.append(f"[Warning: Document paragraph extraction failed: {e}]")

    for table_index, table in enumerate(doc.tables, start=1):
        append_docx_table(text_parts, f"--- Document Table {table_index} ---", table)

    for section_index, section in enumerate(doc.sections, start=1):
        text_parts.append(f"--- Section {section_index} Header ---")
        try:
            for paragraph in section.header.paragraphs:
                paragraph_text = clean_text(paragraph.text)
                if paragraph_text:
                    text_parts.append(paragraph_text)
                _place_images(paragraph, header_rid_map)
        except Exception as e:
            text_parts.append(
                f"[Warning: Header extraction failed for section {section_index}: {e}]"
            )

        for table_index, table in enumerate(section.header.tables, start=1):
            append_docx_table(
                text_parts,
                f"--- Section {section_index} Header Table {table_index} ---",
                table,
            )

        text_parts.append(f"--- Section {section_index} Footer ---")
        try:
            for paragraph in section.footer.paragraphs:
                paragraph_text = clean_text(paragraph.text)
                if paragraph_text:
                    text_parts.append(paragraph_text)
                _place_images(paragraph, footer_rid_map)
        except Exception as e:
            text_parts.append(
                f"[Warning: Footer extraction failed for section {section_index}: {e}]"
            )

        for table_index, table in enumerate(section.footer.tables, start=1):
            append_docx_table(
                text_parts,
                f"--- Section {section_index} Footer Table {table_index} ---",
                table,
            )

    # Herhangi bir nedenle satır içine yerleştirilemeyen görseller (örn. tablo hücresindeki)
    remaining = {
        path
        for rid_map in rid_maps.values()
        for rid, path in rid_map.items()
        if rid not in placed_rids
    }
    if remaining:
        text_parts.append("--- Extracted Images ---")
        for path in sorted(remaining):
            inline_counter[0] += 1
            text_parts.append(
                format_image_marker(path, image_number=inline_counter[0], source="docx")
            )

    return "\n\n".join(part for part in text_parts if part and part.strip())


def extract_text(file_path):
    """
    Dosya uzantısına göre uygun text extraction fonksiyonunu çağırır.
    """
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".txt":
        return extract_text_from_txt(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    raise ValueError(f"Unsupported file type: {extension}")
